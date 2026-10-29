"""Course-template report builder. Run with Python, python-docx and reportlab."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import json,hashlib,subprocess,math
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor,black,white
ROOT=Path(__file__).resolve().parents[2];TMP=ROOT/'tmp/documents';OUT=ROOT/'deliverables/report'
TMP.mkdir(parents=True,exist_ok=True);OUT.mkdir(parents=True,exist_ok=True)
REF=ROOT/'brief/source-materials/Report template.docx'
assert hashlib.sha256(REF.read_bytes()).hexdigest() in (TMP/'artifact.md').read_text()
refs=json.loads((ROOT/'docs/references-verified.json').read_text(encoding='utf-8'))
bench=json.loads((ROOT/'evidence/benchmarks/summary.json').read_text())
POP=next((ROOT/'.tools/poppler').rglob('pdftoppm.exe'));FIG=ROOT/'evidence/figures';FIG.mkdir(exist_ok=True)
def diagram(name):
 p=FIG/(name+'.pdf');c=canvas.Canvas(str(p),pagesize=(720,320));c.setFillColor(white);c.rect(0,0,720,320,stroke=0,fill=1);c.setStrokeColor(HexColor('#293640'));c.setFillColor(black);c.setLineWidth(1.5)
 def box(x,y,w,h,t):
  c.setFillColor(HexColor('#f0f2f3'));c.rect(x,y,w,h,stroke=1,fill=1);c.setFillColor(black);c.setFont('Helvetica',15)
  lines=t.split('\n')
  for i,line in enumerate(lines):c.drawCentredString(x+w/2,y+h/2+(len(lines)-1)*9-i*18-5,line)
 def arrow(x,y,x2,y2):
  c.line(x,y,x2,y2);a=math.atan2(y2-y,x2-x);d=8
  for z in [-.45,.45]:c.line(x2,y2,x2-d*math.cos(a+z),y2-d*math.sin(a+z))
 if name=='architecture':
  box(15,230,145,65,'GLFW input');box(210,230,145,65,'Game\n240 Hz');box(405,230,145,65,'World + rig');arrow(160,263,210,263);arrow(355,263,405,263)
  box(15,130,145,65,'Native assets');box(210,130,145,65,'Scene graph');box(405,130,145,65,'Renderer');arrow(160,163,210,163);arrow(355,163,405,163);arrow(477,230,477,195)
  box(15,25,145,65,'Moon depth');box(210,25,145,65,'CCTV FBO');box(405,25,145,65,'Main HDR');box(590,25,115,65,'Post + UI');arrow(477,130,477,90);arrow(160,58,210,58);arrow(355,58,405,58);arrow(550,58,590,58)
 elif name=='local-box-test':
  c.setFont('Helvetica-Bold',18);c.drawString(30,286,'World space');c.drawString(440,286,'Beam local space')
  c.saveState();c.translate(180,150);c.rotate(22);c.setFillColor(HexColor('#e6b6aa'));c.rect(-140,-16,280,32,fill=1);c.restoreState();c.setDash(5,4);c.rect(44,82,272,136,stroke=1,fill=0);c.setDash()
  c.setFillColor(HexColor('#1a5b83'));c.circle(68,201,5,stroke=0,fill=1);c.setFillColor(black);c.setFont('Helvetica',14);c.drawString(26,237,'Inside world AABB, outside beam');arrow(325,150,410,150);c.drawCentredString(370,179,'M^-1')
  c.setFillColor(HexColor('#e6b6aa'));c.rect(445,134,240,32,fill=1);c.setFillColor(black);c.setFont('Helvetica',16);c.drawString(438,95,'q = M^-1 p');c.drawString(438,65,'|qx| <= hx, |qy| <= hy, |qz| <= hz');c.setFont('Helvetica',14);c.drawString(30,22,'Use the local box, not the larger dashed world AABB.')
 else:
  c.setFont('Helvetica-Bold',18);c.drawString(30,290,'Whole-frame performance at 1280 x 720')
  for i,b in enumerate(bench['runs']):
   y=225-i*70;c.setFont('Helvetica',15);c.drawString(25,y+6,f'Run {i+1}')
   for val,dy,color in [(b['average_fps'],12,'#345e71'),(b['one_percent_low_fps'],-12,'#b18143')]:
    c.setFillColor(HexColor(color));c.rect(100,y+dy,val/3.2,16,stroke=0,fill=1);c.setFillColor(black);c.drawString(110+val/3.2,y+dy+2,f'{val:.1f}')
  c.setFont('Helvetica',14);c.drawString(100,24,'Blue: average FPS     Brown: 1% low FPS     Uncapped; local RTX 5060')
 c.save();subprocess.run([str(POP),'-singlefile','-scale-to','1800','-png',str(p),str(FIG/name)],check=True,capture_output=True)
for name in ['architecture','local-box-test','performance']:diagram(name)
doc=Document(REF);body=doc._element.body
for el in list(body):
 if el.tag!=qn('w:sectPr'):body.remove(el)
for name,size,bold in [('Normal',14,False),('Title',24,True),('Heading 1',18,True),('Heading 2',15,True),('Caption',12,False)]:
 if name not in doc.styles:
  from docx.enum.style import WD_STYLE_TYPE
  doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.08
 st.paragraph_format.first_line_indent=Pt(0);st.paragraph_format.left_indent=Pt(0);st.paragraph_format.right_indent=Pt(0)
 fonts=st.element.get_or_add_rPr().rFonts
 for attr in list(fonts.attrib):
  if 'Theme' in attr:del fonts.attrib[attr]
 if 'Heading' in name:
  st.paragraph_format.keep_with_next=True
  num=st.element.get_or_add_pPr().find(qn('w:numPr'))
  if num is not None:num.getparent().remove(num)
section=doc.sections[0];footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
page_no=0;toc=[];fulltext=[];fig_num=0
def page(title,level=1,listing=True):
 global page_no
 if page_no:doc.add_page_break()
 page_no+=1
 if title:
  doc.add_heading(title,level=level)
  if listing:toc.append((title,page_no))
def para(text,style=None):
 fulltext.append(text);p=doc.add_paragraph(text,style=style);p.paragraph_format.widow_control=True;return p
def h(text):doc.add_heading(text,2)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,width in enumerate(widths):t.columns[i].width=Inches(width)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=str(x)
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row):cells[i].text=str(x)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;props=c._tc.get_or_add_tcPr();bord=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bord.append(e)
   props.append(bord);mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
   props.append(mar)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7ECF0');props.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.0
    for run in p.runs:run.font.size=Pt(12);run.font.bold=ri==0
 para('')
def figure(file,caption,width=6.45):
 global fig_num
 fig_num+=1;p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3);p.add_run().add_picture(str(file),width=Inches(width));para(f'Figure {fig_num}. {caption}','Caption')
def pair(a,b,caption):
 global fig_num
 fig_num+=1
 t=doc.add_table(rows=1,cols=2);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,img in enumerate([a,b]):
  t.columns[i].width=Inches(3.22);cell=t.cell(0,i);cell.width=Inches(3.22)
  p=cell.paragraphs[0];p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(0);p.add_run().add_picture(str(img),width=Inches(3.0))
 para(f'Figure {fig_num}. {caption}','Caption')
def review(ids):
 for i in ids:
  a=refs[i-1];authors=a['authors'];label=authors[0].split()[-1]+(' et al.' if len(authors)>2 else ' and '+authors[1].split()[-1] if len(authors)==2 else '')
  para(f"{label} [{i}]. {a['method_summary']} {a['project_comparison']}")
S=ROOT/'evidence/screenshots'
exec(compile((Path(__file__).parent/'report_content.py').read_text(encoding='utf-8'),'report_content.py','exec'))
anchor=toc_anchor
for title,num in toc:
 el=OxmlElement('w:p');anchor.addnext(el);p=Paragraph(el,doc._body);p.text=f'{title}  . . . . . .  {num}';p.paragraph_format.space_after=Pt(7)
 for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(14)
 anchor=el
doc.core_properties.author='Nurbek Yerbulekov'
doc.core_properties.last_modified_by='Nurbek Yerbulekov'
doc.core_properties.title='Laser Vault Heist'
candidate=TMP/'report-candidate.docx';doc.save(candidate)
editable={'docProps/core.xml','word/document.xml','word/_rels/document.xml.rels','word/styles.xml','[Content_Types].xml'}
final=OUT/'Laser_Vault_Heist_Report.docx'
with ZipFile(REF) as source,ZipFile(candidate) as edited,ZipFile(final,'w',ZIP_DEFLATED) as target:
 for name in source.namelist():target.writestr(name,edited.read(name) if name in editable else source.read(name))
 for name in edited.namelist():
  if name not in source.namelist():target.writestr(name,edited.read(name))
with ZipFile(REF) as source,ZipFile(final) as output:
 unchanged=[n for n in source.namelist() if n not in editable];assert all(source.read(n)==output.read(n) for n in unchanged)
(ROOT/'evidence/report-structure.json').write_text(json.dumps({'expected_pages':page_no,'contents':toc,'reference_sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),'preserved_parts':unchanged,'edited_parts':sorted(editable)},indent=2))
(TMP/'report-prose.txt').write_text('\n\n'.join(fulltext),encoding='utf-8')
print('Created',final,'expected pages',page_no,'preserved opaque parts',len(unchanged))
